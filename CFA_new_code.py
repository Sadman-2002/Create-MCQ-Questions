from flask import Flask, request, jsonify, render_template_string
from werkzeug.utils import secure_filename
import os
import random

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024
app.config['UPLOAD_FOLDER'] = 'uploads'

os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# OpenAI setup
USE_AI = False
client = None
try:
    from openai import OpenAI
    api_key = os.getenv("OPENAI_API_KEY", "YOUR-KEY-HERE")
    if api_key and api_key != "YOUR-KEY-HERE":
        client = OpenAI(api_key=api_key)
        USE_AI = True
        print("✅ OpenAI connected!")
except:
    print("⚠️ No OpenAI, using simple generation")

ALLOWED_EXTENSIONS = {'pdf', 'doc', 'docx', 'ppt', 'pptx', 'xlsx', 'xls', 'txt'}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_from_pdf(file_path):
    try:
        import PyPDF2
        text = ""
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
        return text
    except Exception as e:
        print(f"PDF error: {e}")
        return ""

def extract_text_from_docx(file_path):
    try:
        import docx
        doc = docx.Document(file_path)
        return "\n".join([paragraph.text for paragraph in doc.paragraphs])
    except Exception as e:
        print(f"DOCX error: {e}")
        return ""

def extract_text_from_pptx(file_path):
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        text = ""
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
        return text
    except Exception as e:
        print(f"PPTX error: {e}")
        return ""

def extract_text_from_excel(file_path):
    try:
        import openpyxl
        workbook = openpyxl.load_workbook(file_path)
        text = ""
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = ' '.join([str(cell) for cell in row if cell is not None])
                text += row_text + "\n"
        return text
    except Exception as e:
        print(f"Excel error: {e}")
        return ""

def extract_text_from_txt(file_path):
    try:
        with open(file_path, 'r', encoding='utf-8') as file:
            return file.read()
    except:
        try:
            with open(file_path, 'r', encoding='latin-1') as file:
                return file.read()
        except Exception as e:
            print(f"TXT error: {e}")
            return ""

def extract_text_from_file(file_path, filename):
    extension = filename.rsplit('.', 1)[1].lower()
    
    if extension == 'pdf':
        return extract_text_from_pdf(file_path)
    elif extension in ['doc', 'docx']:
        return extract_text_from_docx(file_path)
    elif extension in ['ppt', 'pptx']:
        return extract_text_from_pptx(file_path)
    elif extension in ['xls', 'xlsx']:
        return extract_text_from_excel(file_path)
    elif extension == 'txt':
        return extract_text_from_txt(file_path)
    else:
        return ""

def generate_simple_mcqs(text, num_questions=10):
    """Generate simple MCQs - EXACT OLD CODE"""
    if len(text) < 100:
        return []
    
    sentences = [s.strip() + '.' for s in text.split('.') if len(s.strip()) > 30]
    
    if len(sentences) < 3:
        return []
    
    questions = []
    num_questions = min(num_questions, len(sentences), 15)
    
    selected_sentences = random.sample(sentences, min(num_questions * 2, len(sentences)))
    
    for sent in selected_sentences:
        if len(questions) >= num_questions:
            break
            
        words = sent.split()
        meaningful_words = [w.strip('.,!?;:()[]{}') for w in words 
                          if len(w) > 4 and w.replace('_', '').isalpha()]
        
        if len(meaningful_words) < 3:
            continue
        
        blank_word = random.choice(meaningful_words)
        question_text = sent.replace(blank_word, "______", 1)
        
        wrong_options = [w for w in meaningful_words if w.lower() != blank_word.lower()]
        random.shuffle(wrong_options)
        wrong_options = wrong_options[:3]
        
        while len(wrong_options) < 3:
            wrong_options.append(f"Alternative {len(wrong_options) + 1}")
        
        all_options = [blank_word] + wrong_options
        random.shuffle(all_options)
        correct_index = all_options.index(blank_word)
        
        questions.append({
            "question": f"Complete the sentence: {question_text}",
            "options": all_options,
            "correct_index": correct_index,
            "explanation": f"The correct answer is '{blank_word}' from the source text: {sent[:100]}..."
        })
    
    return questions

HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>MCQ Generator - DEBUGGING</title>
    <style>
        * { margin: 0; padding: 0; box-sizing: border-box; }
        body {
            background: linear-gradient(135deg, #1e3c72 0%, #2a5298 100%);
            display: flex;
            justify-content: center;
            align-items: center;
            min-height: 100vh;
            font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
            padding: 20px;
        }
        #mainContainer {
            background: white;
            border-radius: 20px;
            box-shadow: 0 20px 60px rgba(0,0,0,0.3);
            max-width: 900px;
            width: 100%;
            padding: 40px;
        }
        h1 {
            color: #1e3c72;
            text-align: center;
            margin-bottom: 10px;
            font-size: 2.5em;
        }
        .subtitle {
            text-align: center;
            color: #666;
            margin-bottom: 30px;
            font-size: 1.1em;
        }
        .upload-area {
            border: 3px dashed #1e3c72;
            border-radius: 15px;
            padding: 60px 40px;
            text-align: center;
            background: #f8f9ff;
            transition: all 0.3s;
            cursor: pointer;
            margin: 30px 0;
        }
        .upload-area:hover { background: #e8ebff; transform: scale(1.02); }
        .upload-icon { font-size: 4em; margin-bottom: 20px; }
        #fileInput { display: none; }
        .btn {
            background: linear-gradient(135deg, #1e3c72 0%, #2a5298 100%);
            color: white;
            padding: 15px 30px;
            border: none;
            border-radius: 10px;
            font-size: 1.1em;
            cursor: pointer;
            transition: transform 0.3s;
            margin: 10px;
        }
        .btn:hover { transform: scale(1.05); }
        .btn:disabled { opacity: 0.5; cursor: not-allowed; }
        .center { text-align: center; }
        .error-message {
            background: #ffebee;
            color: #c62828;
            padding: 15px;
            border-radius: 10px;
            margin: 20px 0;
            display: none;
        }
        .debug-info {
            background: #e3f2fd;
            border: 2px solid #2196f3;
            padding: 15px;
            border-radius: 10px;
            margin: 20px 0;
            font-family: monospace;
            font-size: 0.9em;
        }
        #loadingScreen, #gameScreen, #resultScreen { display: none; }
        .loader {
            border: 5px solid #f3f3f3;
            border-top: 5px solid #1e3c72;
            border-radius: 50%;
            width: 60px;
            height: 60px;
            animation: spin 1s linear infinite;
            margin: 40px auto;
        }
        @keyframes spin {
            0% { transform: rotate(0deg); }
            100% { transform: rotate(360deg); }
        }
        #questionCard {
            background: #f8f9ff;
            padding: 30px;
            border-radius: 15px;
            margin: 20px 0;
        }
        #question { font-size: 1.3em; color: #333; margin-bottom: 25px; line-height: 1.6; }
        .answer-btn {
            display: block;
            width: 100%;
            padding: 15px;
            margin: 10px 0;
            background: white;
            border: 2px solid #1e3c72;
            border-radius: 10px;
            cursor: pointer;
            font-size: 1.1em;
            text-align: left;
            transition: all 0.3s;
        }
        .answer-btn:hover:not(:disabled) {
            background: #1e3c72;
            color: white;
            transform: translateX(5px);
        }
        .answer-btn.correct { background: #4caf50; color: white; border-color: #4caf50; }
        .answer-btn.incorrect { background: #f44336; color: white; border-color: #f44336; }
        .answer-btn:disabled { opacity: 0.7; cursor: not-allowed; }
        #feedback {
            margin-top: 20px;
            padding: 20px;
            border-radius: 10px;
            display: none;
        }
        #feedback.show { display: block; }
        #feedback.correct { background: #e8f5e9; color: #2e7d32; border: 2px solid #4caf50; }
        #feedback.incorrect { background: #ffebee; color: #c62828; border: 2px solid #f44336; }
        .stats {
            display: flex;
            justify-content: space-around;
            margin: 20px 0;
            text-align: center;
        }
        .stat-value { font-size: 2em; font-weight: bold; color: #1e3c72; }
        .stat-label { color: #666; font-size: 0.9em; }
        .progress-bar {
            height: 10px;
            background: #e0e0e0;
            border-radius: 5px;
            overflow: hidden;
            margin: 20px 0;
        }
        .progress-fill {
            height: 100%;
            background: linear-gradient(90deg, #1e3c72 0%, #2a5298 100%);
            transition: width 0.3s;
        }
    </style>
</head>
<body>
    <div id="mainContainer">
        <h1>🔧 MCQ Generator - DEBUG MODE</h1>
        <p class="subtitle">Simplified version to find the problem</p>

        <div id="uploadScreen">
            <div class="debug-info">
                <strong>🔍 DEBUG INFO:</strong><br>
                Status: <span id="debugStatus">Ready to upload</span><br>
                File: <span id="debugFile">None</span><br>
                Size: <span id="debugSize">0 bytes</span><br>
                Last Error: <span id="debugError">None</span>
            </div>

            <div class="upload-area" onclick="document.getElementById('fileInput').click()">
                <div class="upload-icon">📄</div>
                <div id="uploadText">Click to upload</div>
                <div style="color: #666; font-size: 0.9em; margin-top: 10px;">
                    PDF, Word, PowerPoint, Excel, TXT (max 16MB)
                </div>
                <input type="file" id="fileInput" accept=".pdf,.doc,.docx,.ppt,.pptx,.xls,.xlsx,.txt">
            </div>
            
            <div class="center">
                <label style="font-size: 1.1em;">Questions: <span id="numValue">10</span></label><br>
                <input type="range" id="numQuestions" min="5" max="15" value="10" style="width: 300px; margin: 15px 0;">
            </div>

            <div class="center">
                <button class="btn" id="uploadBtn" disabled>🚀 Generate Questions</button>
            </div>

            <div id="errorMsg" class="error-message"></div>
        </div>

        <div id="loadingScreen">
            <div class="loader"></div>
            <div style="text-align: center; color: #1e3c72; font-size: 1.2em; margin: 20px;">
                Processing...<br>
                <span style="font-size: 0.9em; color: #666;">Please wait</span>
            </div>
        </div>

        <div id="gameScreen">
            <div class="progress-bar">
                <div class="progress-fill" id="progressBar" style="width: 0%"></div>
            </div>
            <div class="stats">
                <div class="stat-item">
                    <div class="stat-value" id="currentQ">1</div>
                    <div class="stat-label">Question</div>
                </div>
                <div class="stat-item">
                    <div class="stat-value" id="score">0</div>
                    <div class="stat-label">Score</div>
                </div>
                <div class="stat-item">
                    <div class="stat-value" id="correct">0</div>
                    <div class="stat-label">Correct</div>
                </div>
            </div>
            <div id="questionCard">
                <div id="question"></div>
                <div id="answers"></div>
            </div>
            <div id="feedback">
                <div id="feedbackText"></div>
                <div id="explanation" style="margin-top: 15px; padding: 15px; background: white; border-radius: 8px;"></div>
            </div>
            <div class="center">
                <button class="btn" id="nextBtn" style="display: none;">Next Question →</button>
                <button class="btn" onclick="location.reload()">Start Over</button>
            </div>
        </div>

        <div id="resultScreen">
            <h2 style="text-align: center; margin: 30px 0;">🎉 Quiz Complete!</h2>
            <div class="stats">
                <div class="stat-item">
                    <div class="stat-value" id="finalScore">0</div>
                    <div class="stat-label">Score</div>
                </div>
                <div class="stat-item">
                    <div class="stat-value" id="finalCorrect">0/0</div>
                    <div class="stat-label">Correct</div>
                </div>
                <div class="stat-item">
                    <div class="stat-value" id="finalPercent">0%</div>
                    <div class="stat-label">Accuracy</div>
                </div>
            </div>
            <div id="resultMessage" style="text-align: center; font-size: 1.3em; margin: 30px 0; padding: 25px; background: #f8f9ff; border-radius: 10px;"></div>
            <div class="center">
                <button class="btn" onclick="location.reload()">Upload New Document</button>
            </div>
        </div>
    </div>

    <script>
        let selectedFile = null;
        let questions = [];
        let currentIndex = 0;
        let score = 0;
        let correctCount = 0;

        function updateDebug(status, file = null, error = null) {
            document.getElementById('debugStatus').textContent = status;
            if (file) {
                document.getElementById('debugFile').textContent = file.name;
                document.getElementById('debugSize').textContent = (file.size / 1024).toFixed(2) + ' KB';
            }
            if (error) {
                document.getElementById('debugError').textContent = error;
            }
            console.log('DEBUG:', status, file, error);
        }

        document.getElementById('numQuestions').oninput = function() {
            document.getElementById('numValue').textContent = this.value;
        };

        document.getElementById('fileInput').onchange = function(e) {
            selectedFile = e.target.files[0];
            if (selectedFile) {
                updateDebug('File selected', selectedFile);
                document.getElementById('uploadText').textContent = '✅ ' + selectedFile.name;
                document.getElementById('uploadBtn').disabled = false;
            }
        };

        const uploadArea = document.querySelector('.upload-area');
        uploadArea.addEventListener('dragover', (e) => {
            e.preventDefault();
            uploadArea.style.background = '#e8ebff';
        });
        uploadArea.addEventListener('dragleave', () => {
            uploadArea.style.background = '#f8f9ff';
        });
        uploadArea.addEventListener('drop', (e) => {
            e.preventDefault();
            uploadArea.style.background = '#f8f9ff';
            const files = e.dataTransfer.files;
            if (files.length > 0) {
                selectedFile = files[0];
                updateDebug('File dropped', selectedFile);
                document.getElementById('uploadText').textContent = '✅ ' + selectedFile.name;
                document.getElementById('uploadBtn').disabled = false;
            }
        });

        document.getElementById('uploadBtn').onclick = async function() {
            if (!selectedFile) {
                updateDebug('ERROR: No file', null, 'No file selected');
                return;
            }

            const formData = new FormData();
            formData.append('file', selectedFile);
            formData.append('num_questions', document.getElementById('numQuestions').value);

            updateDebug('Uploading file...');
            showScreen('loadingScreen');

            try {
                updateDebug('Sending request to server...');
                
                const response = await fetch('/upload', {
                    method: 'POST',
                    body: formData
                });

                updateDebug('Received response, parsing...');
                const data = await response.json();

                console.log('Server response:', data);

                if (data.error) {
                    updateDebug('ERROR from server', null, data.error);
                    showError(data.error);
                    showScreen('uploadScreen');
                    return;
                }

                updateDebug('Success! Questions received: ' + data.questions.length);
                questions = data.questions;
                currentIndex = 0;
                score = 0;
                correctCount = 0;
                showScreen('gameScreen');
                loadQuestion();

            } catch (error) {
                updateDebug('FETCH ERROR', null, error.message);
                console.error('Full error:', error);
                showError('Error: ' + error.message);
                showScreen('uploadScreen');
            }
        };

        function loadQuestion() {
            if (currentIndex >= questions.length) {
                showResults();
                return;
            }

            const q = questions[currentIndex];
            document.getElementById('currentQ').textContent = currentIndex + 1;
            document.getElementById('progressBar').style.width = ((currentIndex / questions.length) * 100) + '%';
            document.getElementById('question').textContent = q.question;

            const answersDiv = document.getElementById('answers');
            answersDiv.innerHTML = '';

            q.options.forEach((opt, i) => {
                const btn = document.createElement('button');
                btn.className = 'answer-btn';
                btn.textContent = String.fromCharCode(65 + i) + ') ' + opt;
                btn.onclick = () => checkAnswer(i);
                answersDiv.appendChild(btn);
            });

            document.getElementById('feedback').className = '';
            document.getElementById('feedback').style.display = 'none';
            document.getElementById('nextBtn').style.display = 'none';
        }

        function checkAnswer(selected) {
            const q = questions[currentIndex];
            const buttons = document.querySelectorAll('.answer-btn');

            buttons.forEach((btn, i) => {
                btn.disabled = true;
                if (i === q.correct_index) {
                    btn.classList.add('correct');
                } else if (i === selected) {
                    btn.classList.add('incorrect');
                }
            });

            const isCorrect = selected === q.correct_index;
            
            if (isCorrect) {
                score += 100;
                correctCount++;
                document.getElementById('feedbackText').textContent = '✅ Correct!';
                document.getElementById('feedback').className = 'correct show';
            } else {
                document.getElementById('feedbackText').textContent = '❌ Incorrect. Answer: ' + String.fromCharCode(65 + q.correct_index);
                document.getElementById('feedback').className = 'incorrect show';
            }

            document.getElementById('explanation').innerHTML = '<strong>Explanation:</strong><br>' + q.explanation;
            document.getElementById('score').textContent = score;
            document.getElementById('correct').textContent = correctCount;
            document.getElementById('nextBtn').style.display = 'inline-block';
        }

        document.getElementById('nextBtn').onclick = function() {
            currentIndex++;
            loadQuestion();
        };

        function showResults() {
            showScreen('resultScreen');
            const percent = Math.round((correctCount / questions.length) * 100);
            
            document.getElementById('finalScore').textContent = score;
            document.getElementById('finalCorrect').textContent = correctCount + '/' + questions.length;
            document.getElementById('finalPercent').textContent = percent + '%';

            let msg = percent >= 70 ? '🎉 Great job!' : '📚 Keep studying!';
            document.getElementById('resultMessage').textContent = msg;
        }

        function showScreen(screen) {
            ['uploadScreen', 'loadingScreen', 'gameScreen', 'resultScreen'].forEach(s => {
                document.getElementById(s).style.display = 'none';
            });
            document.getElementById(screen).style.display = 'block';
        }

        function showError(msg) {
            const errorDiv = document.getElementById('errorMsg');
            errorDiv.textContent = '❌ ' + msg;
            errorDiv.style.display = 'block';
            setTimeout(() => errorDiv.style.display = 'none', 8000);
        }
    </script>
</body>
</html>
"""

@app.route('/')
def index():
    return render_template_string(HTML_TEMPLATE)

@app.route('/upload', methods=['POST'])
def upload_file():
    print("\n" + "="*60)
    print("📤 UPLOAD REQUEST RECEIVED")
    print("="*60)
    
    if 'file' not in request.files:
        print("❌ ERROR: No file in request")
        return jsonify({'error': 'No file provided'}), 400
    
    file = request.files['file']
    print(f"📄 File object: {file}")
    print(f"📄 Filename: {file.filename}")
    
    if file.filename == '':
        print("❌ ERROR: Empty filename")
        return jsonify({'error': 'No file selected'}), 400
    
    if not allowed_file(file.filename):
        print(f"❌ ERROR: Invalid file type: {file.filename}")
        return jsonify({'error': 'Invalid file type'}), 400
    
    try:
        filename = secure_filename(file.filename)
        file_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        
        print(f"💾 Saving to: {file_path}")
        file.save(file_path)
        print(f"✅ File saved: {os.path.getsize(file_path)} bytes")
        
        print(f"📖 Extracting text...")
        text = extract_text_from_file(file_path, filename)
        print(f"✅ Extracted {len(text)} characters")
        print(f"📝 Preview: {text[:200]}...")
        
        if not text or len(text.strip()) < 50:
            print(f"❌ ERROR: Not enough text (got {len(text)} chars)")
            os.remove(file_path)
            return jsonify({'error': f'Could not extract enough text. Only got {len(text)} characters.'}), 400
        
        num_questions = int(request.form.get('num_questions', 10))
        print(f"🎯 Generating {num_questions} questions...")
        
        mcqs = generate_simple_mcqs(text, num_questions)
        print(f"✅ Generated {len(mcqs)} questions")
        
        os.remove(file_path)
        print(f"🗑️ Cleaned up file")
        
        if not mcqs:
            print(f"❌ ERROR: No questions generated")
            return jsonify({'error': 'Could not generate questions'}), 400
        
        print(f"✅ SUCCESS! Returning {len(mcqs)} questions")
        print("="*60 + "\n")
        
        return jsonify({
            'success': True,
            'questions': mcqs,
            'debug': {
                'text_length': len(text),
                'questions_generated': len(mcqs)
            }
        })
    
    except Exception as e:
        print(f"❌ EXCEPTION: {type(e).__name__}: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({'error': f'Server error: {str(e)}'}), 500

if __name__ == '__main__':
    print("\n" + "="*60)
    print("🚀 STARTING DEBUG VERSION")
    print("="*60)
    print("📂 Open: http://localhost:5000")
    print("🔍 Watch this console for DEBUG messages")
    print("="*60 + "\n")
    app.run(debug=True, port=5000, use_reloader=False)